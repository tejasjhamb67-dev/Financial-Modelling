"""Cell registry - the backbone of genuine cross-sheet linkage.

Every model line is registered as (sheet, metric, period) -> cell address in a
layout pass, so that later, when writing a formula on any sheet, we can emit a
real reference like ``='Income Statement'!F10`` instead of duplicating a hard
number. It also records a lineage entry per cell (data type + source) so the
system can answer "where did this number come from and where does it flow?".

Columns are consistent across sheets: column A holds labels, and each period
occupies the same column on every statement/schedule sheet, which keeps the
workbook navigable.
"""

from __future__ import annotations

from dataclasses import dataclass, field
from typing import Optional

from openpyxl.utils import get_column_letter


@dataclass
class Lineage:
    sheet: str
    metric: str
    period: str
    address: str
    data_type: str
    source_id: Optional[str] = None
    source_document: Optional[str] = None
    page: Optional[int] = None
    formula: Optional[str] = None


class CellRegistry:
    def __init__(self, all_periods: list[str], label_col: int = 1, first_data_col: int = 2):
        self.all_periods = all_periods
        self.label_col = label_col
        self.first_data_col = first_data_col
        self._cells: dict[tuple[str, str, str], str] = {}
        self._rows: dict[tuple[str, str], int] = {}
        self.lineage: list[Lineage] = []

    # -- columns ---------------------------------------------------------
    def period_col_idx(self, period: str) -> int:
        return self.first_data_col + self.all_periods.index(period)

    def period_col_letter(self, period: str) -> str:
        return get_column_letter(self.period_col_idx(period))

    def label_col_letter(self) -> str:
        return get_column_letter(self.label_col)

    # -- registration ----------------------------------------------------
    def set_row(self, sheet: str, key: str, row: int) -> None:
        self._rows[(sheet, key)] = row

    def row(self, sheet: str, key: str) -> Optional[int]:
        return self._rows.get((sheet, key))

    def register(self, sheet: str, metric: str, period: str, row: int) -> str:
        addr = f"{self.period_col_letter(period)}{row}"
        self._cells[(sheet, metric, period)] = addr
        return addr

    def register_addr(self, sheet: str, metric: str, period: str, addr: str) -> None:
        self._cells[(sheet, metric, period)] = addr

    # -- lookups ---------------------------------------------------------
    def addr(self, sheet: str, metric: str, period: str) -> Optional[str]:
        return self._cells.get((sheet, metric, period))

    def has(self, sheet: str, metric: str, period: str) -> bool:
        return (sheet, metric, period) in self._cells

    def ref(self, sheet: str, metric: str, period: str) -> Optional[str]:
        addr = self.addr(sheet, metric, period)
        if addr is None:
            return None
        return f"'{sheet}'!{addr}"

    def local(self, metric_addr: str) -> str:
        return metric_addr

    # -- lineage ---------------------------------------------------------
    def add_lineage(self, entry: Lineage) -> None:
        self.lineage.append(entry)
