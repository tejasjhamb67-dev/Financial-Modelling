"""Helper for custom-layout sheets (valuation, sensitivities, cover, sources, checks).

Provides a thin writer over a worksheet that applies the same colour/formatting
conventions and records lineage, for sheets whose layout doesn't fit the regular
period grid.
"""

from __future__ import annotations

from typing import Optional

from openpyxl.comments import Comment
from openpyxl.utils import get_column_letter

from excel.formatting import styles
from excel.linking.cell_registry import CellRegistry, Lineage
from modules.schemas import DataType


class CustomSheet:
    def __init__(self, ws, registry: CellRegistry, name: str):
        self.ws = ws
        self.registry = registry
        self.name = name
        self.ws.sheet_view.showGridLines = False

    def title(self, text: str, span: int = 8) -> None:
        c = self.ws.cell(row=1, column=1, value=text)
        c.font = styles.header_font()
        for i in range(1, span + 1):
            self.ws.cell(row=1, column=i).fill = styles.header_fill()

    def subtitle(self, text: str, row: int = 2) -> None:
        c = self.ws.cell(row=row, column=1, value=text)
        c.font = styles.base_font(styles._C["subtle_grey"])

    def section(self, row: int, text: str, span: int = 8) -> None:
        c = self.ws.cell(row=row, column=1, value=text)
        c.font = styles.section_font()
        for i in range(1, span + 1):
            self.ws.cell(row=row, column=i).fill = styles.subheader_fill()

    def label(self, row: int, col: int, text: str, bold: bool = False,
              color: Optional[str] = None) -> None:
        c = self.ws.cell(row=row, column=col, value=text)
        c.font = styles.base_font(color, bold=bold)

    def put(self, row: int, col: int, content, data_type: DataType = DataType.DERIVED,
            number_format: str = styles.NF_CURRENCY, comment: Optional[str] = None,
            bold: bool = False, key: Optional[str] = None, period: Optional[str] = None) -> str:
        cell = self.ws.cell(row=row, column=col, value=content)
        is_formula = isinstance(content, str) and content.startswith("=")
        is_link = is_formula and "!" in content
        if data_type == DataType.ASSUMPTION:
            cell.fill = styles.assumption_fill()
        cell.font = styles.base_font(styles.font_color_for(data_type, is_link=is_link), bold=bold)
        cell.number_format = number_format
        cell.alignment = styles.RIGHT
        if comment:
            cell.comment = Comment(comment, "FM Agent")
        addr = f"{get_column_letter(col)}{row}"
        if key and period:
            self.registry.register_addr(self.name, key, period, addr)
        if key and not period:
            self.registry.register_addr(self.name, key, "_", addr)
        self.registry.add_lineage(Lineage(
            sheet=self.name, metric=key or "", period=period or "_", address=addr,
            data_type=data_type.value, formula=content if is_formula else None))
        return addr

    def width(self, col: int, w: float) -> None:
        self.ws.column_dimensions[get_column_letter(col)].width = w
