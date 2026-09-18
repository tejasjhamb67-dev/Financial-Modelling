"""Valuation and Sensitivities sheets (frameworks only).

The agent builds valuation *infrastructure* - a working DCF, a trading-comps
template and (where segments exist) a SOTP scaffold - but never selects the
preferred methodology or multiple, and never states a target price or
conclusion. WACC, terminal growth, target multiples, peer inputs and share
count are analyst inputs (yellow).
"""

from __future__ import annotations

from openpyxl.utils import get_column_letter

from excel.formatting import styles
from excel.workbook_builder.custom import CustomSheet
from modules.model_spec import ModelSpec
from modules.schemas import DataType, SourceDatabase


def fill_valuation(ws, registry, spec: ModelSpec, db: SourceDatabase) -> None:
    cs = CustomSheet(ws, registry, "Valuation")
    cs.title("Valuation - Frameworks", span=8)
    cs.subtitle("Framework only. WACC, terminal growth, target multiple, peers and shares are "
                "ANALYST inputs (yellow). The agent does not select a method or a target price.")
    cs.width(1, 40)
    for c in range(2, 9):
        cs.width(c, 14)

    fps = spec.forecast_periods
    ff = fps[0] if fps else None
    last_hist = spec.historical_periods[-1] if spec.historical_periods else None
    wacc = registry.ref("Assumptions", "assum_wacc", ff)
    tg = registry.ref("Assumptions", "assum_terminal_growth", ff)
    mult = registry.ref("Assumptions", "assum_target_multiple", ff)

    # ---- DCF table ----
    cs.section(4, "Discounted cash flow (FCFF)")
    hdr = 5
    cs.label(hdr, 1, "")
    col_of = {}
    for i, p in enumerate(fps):
        col = 2 + i
        col_of[p] = col
        c = ws.cell(row=hdr, column=col, value=p)
        c.font = styles.section_font()
        c.alignment = styles.CENTER

    def ref_is(m, p):
        return registry.ref("Income Statement", m, p)

    def ref_cf(m, p):
        return registry.ref("Cash Flow", m, p)

    def ref_wc(m, p):
        return registry.ref("Working Capital", m, p)

    rows = {
        "ebit": 6, "tax": 7, "nopat": 8, "dna": 9, "capex": 10,
        "dnwc": 11, "fcff": 12, "n": 13, "df": 14, "pv": 15,
    }
    cs.label(rows["ebit"], 1, "EBIT")
    cs.label(rows["tax"], 1, "(-) Tax on EBIT")
    cs.label(rows["nopat"], 1, "NOPAT", bold=True)
    cs.label(rows["dna"], 1, "(+) D&A")
    cs.label(rows["capex"], 1, "(-) Capex")
    cs.label(rows["dnwc"], 1, "(-) Change in NWC")
    cs.label(rows["fcff"], 1, "= Free cash flow to firm", bold=True)
    cs.label(rows["n"], 1, "Discount period")
    cs.label(rows["df"], 1, "Discount factor")
    cs.label(rows["pv"], 1, "PV of FCFF", bold=True)

    for i, p in enumerate(fps):
        col = col_of[p]
        L = get_column_letter(col)
        ebit = ref_is("ebit", p)
        cs.put(rows["ebit"], col, f"={ebit}" if ebit else None, DataType.LINKED)
        tr = registry.ref("Assumptions", "assum_tax_rate", p)
        cs.put(rows["tax"], col, f"={L}{rows['ebit']}*{tr}", DataType.DERIVED)
        cs.put(rows["nopat"], col, f"={L}{rows['ebit']}-{L}{rows['tax']}", DataType.DERIVED, bold=True)
        dna = ref_is("depreciation", p)
        cs.put(rows["dna"], col, f"={dna}" if dna else 0, DataType.LINKED if dna else DataType.DERIVED)
        capex = ref_cf("capex", p)
        cs.put(rows["capex"], col, f"={capex}" if capex else 0, DataType.LINKED if capex else DataType.DERIVED)
        # change in NWC
        prev_p = spec.all_periods[spec.all_periods.index(p) - 1]
        nwc_t = ref_wc("net_working_capital", p)
        nwc_p = ref_wc("net_working_capital", prev_p)
        if nwc_t and nwc_p:
            cs.put(rows["dnwc"], col, f"={nwc_t}-{nwc_p}", DataType.DERIVED,
                   comment="Increase in NWC is a cash outflow")
        else:
            cs.put(rows["dnwc"], col, 0, DataType.DERIVED)
        cs.put(rows["fcff"], col,
               f"={L}{rows['nopat']}+{L}{rows['dna']}-{L}{rows['capex']}-{L}{rows['dnwc']}",
               DataType.DERIVED, bold=True, key="dcf_fcff", period=p)
        cs.put(rows["n"], col, i + 1, DataType.DERIVED, number_format=styles.NF_RATIO)
        cs.put(rows["df"], col, f"=1/(1+{wacc})^{L}{rows['n']}", DataType.DERIVED,
               number_format=styles.NF_RATIO)
        cs.put(rows["pv"], col, f"={L}{rows['fcff']}*{L}{rows['df']}", DataType.DERIVED, bold=True)

    # ---- DCF summary ----
    s = 17
    pv_range_terms = "+".join(f"{get_column_letter(col_of[p])}{rows['pv']}" for p in fps)
    fcff_last = f"{get_column_letter(col_of[fps[-1]])}{rows['fcff']}" if fps else "0"
    df_last = f"{get_column_letter(col_of[fps[-1]])}{rows['df']}" if fps else "1"
    netdebt = registry.ref("3-Statement Model", "net_debt", last_hist)

    cs.section(s, "DCF summary")
    cs.label(s + 1, 1, "Sum PV of explicit FCFF")
    a_sumpv = cs.put(s + 1, 2, f"={pv_range_terms}", DataType.DERIVED, key="dcf_sumpv")
    cs.label(s + 2, 1, "Terminal value (Gordon growth)")
    a_tv = cs.put(s + 2, 2, f"={fcff_last}*(1+{tg})/({wacc}-{tg})", DataType.DERIVED,
                  comment="TV = FCFF_n*(1+g)/(WACC-g)", key="dcf_tv")
    cs.label(s + 3, 1, "PV of terminal value")
    a_pvtv = cs.put(s + 3, 2, f"=B{s+2}*{df_last}", DataType.DERIVED, key="dcf_pvtv")
    cs.label(s + 4, 1, "Enterprise value", bold=True)
    a_ev = cs.put(s + 4, 2, f"=B{s+1}+B{s+3}", DataType.DERIVED, bold=True, key="dcf_ev")
    cs.label(s + 5, 1, "(-) Net debt (current)")
    cs.put(s + 5, 2, f"={netdebt}" if netdebt else 0, DataType.LINKED if netdebt else DataType.DERIVED,
           key="dcf_netdebt")
    cs.label(s + 6, 1, "Equity value", bold=True)
    cs.put(s + 6, 2, f"=B{s+4}-B{s+5}", DataType.DERIVED, bold=True, key="dcf_equity")
    cs.label(s + 7, 1, "Shares outstanding")
    cs.put(s + 7, 2, _shares_seed(db), DataType.ASSUMPTION, number_format=styles.NF_CURRENCY,
           comment="Analyst input: diluted shares outstanding.", key="shares")
    cs.label(s + 8, 1, "Implied value per share (DCF)", bold=True)
    cs.put(s + 8, 2, f"=B{s+6}/B{s+7}", DataType.DERIVED, bold=True, number_format=styles.NF_RATIO,
           key="dcf_value_per_share")
    cs.label(s + 9, 1, "WACC / terminal growth (from Assumptions)")
    cs.put(s + 9, 2, f"={wacc}", DataType.LINKED, number_format=styles.NF_PERCENT)
    cs.put(s + 9, 3, f"={tg}", DataType.LINKED, number_format=styles.NF_PERCENT)

    # ---- EV/EBITDA cross-check ----
    e = s + 11
    cs.section(e, "EV/EBITDA cross-check (analyst target multiple)")
    ebitda_last = registry.ref("Income Statement", "ebitda", last_hist)
    cs.label(e + 1, 1, "Target EV/EBITDA (x)")
    cs.put(e + 1, 2, f"={mult}", DataType.LINKED, number_format=styles.NF_MULTIPLE)
    cs.label(e + 2, 1, f"EBITDA ({last_hist})")
    cs.put(e + 2, 2, f"={ebitda_last}" if ebitda_last else None, DataType.LINKED)
    cs.label(e + 3, 1, "Implied enterprise value")
    cs.put(e + 3, 2, f"=B{e+1}*B{e+2}", DataType.DERIVED)
    cs.label(e + 4, 1, "Implied equity value")
    cs.put(e + 4, 2, f"=B{e+3}-B{s+5}", DataType.DERIVED)
    cs.label(e + 5, 1, "Implied value per share (multiple)", bold=True)
    cs.put(e + 5, 2, f"=B{e+4}/B{s+7}", DataType.DERIVED, bold=True, number_format=styles.NF_RATIO)

    # ---- Trading comparables template ----
    t = e + 7
    cs.section(t, "Trading comparables (analyst to populate peers)")
    headers = ["Peer", "Market cap", "Enterprise value", "Revenue", "EBITDA", "EBIT",
               "PAT", "EV/EBITDA", "P/E"]
    for j, h in enumerate(headers):
        c = ws.cell(row=t + 1, column=1 + j, value=h)
        c.font = styles.section_font()
    for r in range(t + 2, t + 6):
        cs.put(r, 1, "<peer>", DataType.ASSUMPTION, number_format="@")
        for j in range(1, 7):
            cs.put(r, 1 + j, None, DataType.ASSUMPTION)
        # EV/EBITDA = EnterpriseValue(C) / EBITDA(E); P/E = MarketCap(B) / PAT(G)
        cs.put(r, 8, f"=IF(E{r}=0,\"\",C{r}/E{r})", DataType.DERIVED, number_format=styles.NF_MULTIPLE)
        cs.put(r, 9, f"=IF(G{r}=0,\"\",B{r}/G{r})", DataType.DERIVED, number_format=styles.NF_MULTIPLE)
    cs.label(t + 6, 1, "Median EV/EBITDA")
    cs.put(t + 6, 8, f"=IFERROR(MEDIAN(H{t+2}:H{t+5}),\"\")", DataType.DERIVED, number_format=styles.NF_MULTIPLE)

    # ---- SOTP scaffold ----
    if "SOTP" in spec.valuation_methods:
        so = t + 8
        cs.section(so, "Sum-of-the-parts (segment scaffold)")
        cs.label(so + 1, 1, "Segment")
        for j, h in enumerate(["Metric", "Multiple (x)", "Implied value"]):
            ws.cell(row=so + 1, column=2 + j, value=h).font = styles.section_font()
        cs.subtitle("Populate per-segment metric and multiple; agent does not choose multiples.",
                    row=so + 2)


def _shares_seed(db: SourceDatabase):
    # Do not fabricate a share count; use a placeholder the analyst must set.
    return 100.0


def fill_sensitivities(ws, registry, spec: ModelSpec, db: SourceDatabase) -> None:
    cs = CustomSheet(ws, registry, "Sensitivities")
    cs.title("Sensitivities", span=8)
    cs.subtitle("Grids recompute the DCF implied value per share. Axis values recenter on the "
                "Assumptions inputs; analyst controls the ranges.")
    cs.width(1, 26)
    for c in range(2, 9):
        cs.width(c, 13)

    fps = spec.forecast_periods
    if not fps:
        cs.section(4, "Insufficient forecast periods for sensitivities.")
        return
    ff = fps[0]
    wacc = registry.ref("Assumptions", "assum_wacc", ff)
    tg = registry.ref("Assumptions", "assum_terminal_growth", ff)
    shares = registry.ref("Valuation", "shares", "_")
    netdebt = registry.ref("Valuation", "dcf_netdebt", "_")
    fcff_refs = [registry.ref("Valuation", "dcf_fcff", p) for p in fps]

    # ---- WACC (cols) vs terminal growth (rows) ----
    cs.section(4, "Implied value per share: WACC (columns) vs terminal growth (rows)")
    top = 5
    # WACC across columns B..F, offsets
    wacc_offsets = [-0.01, -0.005, 0.0, 0.005, 0.01]
    g_offsets = [-0.01, -0.005, 0.0, 0.005, 0.01]
    cs.label(top, 1, "g \\ WACC", bold=True)
    for j, wo in enumerate(wacc_offsets):
        col = 2 + j
        cs.put(top, col, f"={wacc}{_sign(wo)}", DataType.DERIVED, number_format=styles.NF_PERCENT, bold=True)
    for i, go in enumerate(g_offsets):
        row = top + 1 + i
        cs.put(row, 1, f"={tg}{_sign(go)}", DataType.DERIVED, number_format=styles.NF_PERCENT, bold=True)
        for j, wo in enumerate(wacc_offsets):
            col = 2 + j
            w_ref = f"{get_column_letter(col)}${top}"
            g_ref = f"$A{row}"
            formula = _dcf_value_formula(fcff_refs, w_ref, g_ref, netdebt, shares)
            cs.put(row, col, formula, DataType.DERIVED, number_format=styles.NF_RATIO)

    cs.subtitle("Value per share = [ Σ FCFFt/(1+WACC)^t + TV/(1+WACC)^n - net debt ] / shares, "
                "TV = FCFFn*(1+g)/(WACC-g).", row=top + len(g_offsets) + 2)


def _dcf_value_formula(fcff_refs, w_ref, g_ref, netdebt, shares):
    n = len(fcff_refs)
    pv_terms = [f"{fcff_refs[i]}/(1+{w_ref})^{i+1}" for i in range(n)]
    last = fcff_refs[-1]
    tv = f"({last}*(1+{g_ref})/({w_ref}-{g_ref}))/(1+{w_ref})^{n}"
    ev = "+".join(pv_terms + [tv])
    nd = netdebt or "0"
    sh = shares or "1"
    return f"=(({ev})-{nd})/{sh}"


def _sign(x: float) -> str:
    if x == 0:
        return "+0"
    return (f"+{x}" if x > 0 else f"-{abs(x)}")
