"""Cover, Sources and Checks sheets."""

from __future__ import annotations

from openpyxl.formatting.rule import FormulaRule
from openpyxl.utils import get_column_letter

from excel.formatting import styles
from excel.formulas import fw
from excel.workbook_builder.custom import CustomSheet
from modules.config_loader import validation_config
from modules.model_spec import ModelSpec
from modules.schemas import Consolidation, DataType, SourceDatabase, Statement

_CONSOL = Consolidation.CONSOLIDATED


# ---------------------------------------------------------------------------
# Cover
# ---------------------------------------------------------------------------

def fill_cover(ws, registry, spec: ModelSpec, db: SourceDatabase, build_date: str) -> None:
    cs = CustomSheet(ws, registry, "Cover")
    cs.width(1, 34)
    cs.width(2, 60)
    t = ws.cell(row=1, column=1, value=spec.company)
    t.font = styles.title_font()
    for c in range(1, 6):
        ws.cell(row=1, column=c).fill = styles.header_fill()
    ws.cell(row=2, column=1, value="Financial Model (agent-constructed historicals + "
                                   "analyst-driven forecast)").font = styles.base_font(styles._C["subtle_grey"])

    rows = [
        ("Model built", build_date),
        ("Reporting currency", spec.currency or "Not established"),
        ("Reporting unit", spec.unit or "Not established"),
        ("Consolidation basis", spec.consolidation),
        ("Historical periods", ", ".join(spec.historical_periods) or "None"),
        ("Forecast periods", ", ".join(spec.forecast_periods) or "None"),
        ("Documents used", str(len(db.documents))),
    ]
    r = 4
    cs.section(r, "Model information")
    r += 1
    for label, val in rows:
        cs.label(r, 1, label, bold=True)
        ws.cell(row=r, column=2, value=val).font = styles.base_font()
        r += 1

    r += 1
    cs.section(r, "Scope & responsibility")
    r += 1
    disc = [
        "The agent automates the mechanical, historical and diligence work: finding, extracting,",
        "structuring, normalizing and constructing historical financials, linking the model, and",
        "building forecast/valuation/sensitivity INFRASTRUCTURE plus checks.",
        "The ANALYST owns all forecast assumptions, business outlook, valuation assumptions,",
        "scenarios, the investment thesis, target price and conclusion.",
        "Numbers are never invented: missing items are shown as N/A / Not Found / Requires Review.",
    ]
    for line in disc:
        ws.cell(row=r, column=1, value=line).font = styles.base_font()
        r += 1

    r += 1
    cs.section(r, "Colour legend")
    r += 1
    legend = [
        ("Blue", "Reported / hardcoded historical input", DataType.REPORTED),
        ("Black", "Formula within the same worksheet", DataType.DERIVED),
        ("Green", "Link to another worksheet", DataType.LINKED),
        ("Yellow fill", "User assumption (analyst input)", DataType.ASSUMPTION),
    ]
    for name, desc, dt in legend:
        c = ws.cell(row=r, column=1, value=name)
        c.font = styles.base_font(styles.font_color_for(dt))
        if dt == DataType.ASSUMPTION:
            c.fill = styles.assumption_fill()
        ws.cell(row=r, column=2, value=desc).font = styles.base_font()
        r += 1

    r += 1
    cs.section(r, "Sheet index")
    r += 1
    for s in spec.sheets:
        ws.cell(row=r, column=1, value=s).font = styles.link_font()
        r += 1

    if spec.exceptions:
        r += 1
        cs.section(r, "Analyst review required")
        r += 1
        for exc in spec.exceptions:
            c = ws.cell(row=r, column=1, value="• " + exc)
            c.font = styles.error_font()
            r += 1


# ---------------------------------------------------------------------------
# Sources
# ---------------------------------------------------------------------------

def fill_sources(ws, registry, spec: ModelSpec, db: SourceDatabase) -> None:
    cs = CustomSheet(ws, registry, "Sources")
    cs.title("Source Registry", span=10)
    cs.subtitle("Full provenance for every extracted datapoint. Nothing here is invented.")

    # document inventory
    r = 4
    cs.section(r, "Document inventory")
    r += 1
    dh = ["Doc ID", "Name", "Type", "FY", "Consolidation", "Currency", "Unit", "Pages"]
    for j, h in enumerate(dh):
        ws.cell(row=r, column=1 + j, value=h).font = styles.section_font()
    r += 1
    for d in db.documents:
        vals = [d.doc_id, d.name, d.doc_type, d.fiscal_year, d.consolidation.value,
                d.currency, d.unit, ",".join(str(p) for p in (d.relevant_pages or []))]
        for j, v in enumerate(vals):
            ws.cell(row=r, column=1 + j, value=v).font = styles.base_font()
        r += 1

    r += 1
    cs.section(r, "Datapoints")
    r += 1
    headers = ["Source ID", "Metric", "Statement", "Period", "Value", "Unit", "Type",
               "Document", "Page", "Status", "Notes"]
    for j, h in enumerate(headers):
        ws.cell(row=r, column=1 + j, value=h).font = styles.section_font()
    r += 1
    widths = [26, 22, 10, 8, 14, 12, 12, 26, 6, 18, 40]
    for j, w in enumerate(widths):
        cs.width(1 + j, w)

    for dp in sorted(db.datapoints, key=lambda x: (x.statement.value, x.metric, x.period)):
        row = [dp.source_id, dp.metric, dp.statement.value, dp.period,
               dp.value if dp.value is not None else "Not Found",
               dp.unit, dp.data_type.value, dp.source_document, dp.page,
               dp.status.value, dp.notes]
        for j, v in enumerate(row):
            cell = ws.cell(row=r, column=1 + j, value=v)
            cell.font = styles.base_font()
            if j == 4 and dp.value is not None:
                cell.number_format = styles.NF_CURRENCY
                cell.font = styles.input_font()
            if j == 4 and dp.value is None:
                cell.font = styles.error_font()
        r += 1


# ---------------------------------------------------------------------------
# Checks
# ---------------------------------------------------------------------------

def fill_checks(ws, registry, spec: ModelSpec, db: SourceDatabase) -> None:
    cs = CustomSheet(ws, registry, "Checks")
    cs.title("Model Checks", span=2 + len(spec.all_periods))
    cs.subtitle("Reconciliations. Checks are never forced to pass - genuine discrepancies are shown.")
    cs.width(1, 40)

    vcfg = validation_config()["tolerances"]
    bs_abs = vcfg["balance_sheet_abs"]
    bs_rel = vcfg["balance_sheet_rel"]
    cf_abs = vcfg["cash_flow_abs"]

    periods = spec.all_periods
    col_of = {}
    hdr = 4
    cs.label(hdr, 1, "Check", bold=True)
    for i, p in enumerate(periods):
        col = 2 + i
        col_of[p] = col
        cs.width(col, 12)
        c = ws.cell(row=hdr, column=col, value=p)
        c.font = styles.section_font()
        c.alignment = styles.CENTER

    row = hdr + 1
    status_cells: list[str] = []

    def add_check(label, diff_fn, tol_fn, note=""):
        nonlocal row
        cs.label(row, 1, label + " (difference)")
        diff_row = row
        for p in periods:
            f = diff_fn(p)
            if f:
                cs.put(diff_row, col_of[p], f, DataType.DERIVED)
        row += 1
        cs.label(row, 1, label + " (status)")
        for p in periods:
            d_addr = f"{get_column_letter(col_of[p])}{diff_row}"
            tol = tol_fn(p)
            if tol is None:
                continue
            status = (f'=IF({d_addr}="","N/A",IF(ABS({d_addr})<={tol},"PASS",'
                      f'IF(ABS({d_addr})<={tol}*10,"WARNING","ERROR")))')
            addr = cs.put(row, col_of[p], status, DataType.DERIVED, number_format="@")
            status_cells.append(addr)
        if note:
            row += 1
            ws.cell(row=row, column=1, value="   " + note).font = styles.base_font(styles._C["subtle_grey"])
        row += 2

    # 1. Balance sheet identity
    def bs_diff(p):
        ta = registry.ref("Balance Sheet", "total_assets", p)
        tel = registry.ref("Balance Sheet", "total_equity_and_liabilities", p)
        return f"={ta}-{tel}" if ta and tel else None

    def bs_tol(p):
        ta = registry.ref("Balance Sheet", "total_assets", p)
        return f"MAX({bs_abs},{bs_rel}*ABS({ta}))" if ta else None

    add_check("Balance sheet: Assets - (Equity + Liabilities)", bs_diff, bs_tol,
              "Expected 0. Historical gaps indicate reported totals vs reconstructed components.")

    # 2. Cash flow to balance sheet cash
    def cf_diff(p):
        cc = registry.ref("Cash Flow", "closing_cash", p)
        cash = registry.ref("Balance Sheet", "cash", p)
        return f"={cc}-{cash}" if cc and cash else None
    add_check("Cash flow closing cash vs Balance Sheet cash", cf_diff, lambda p: cf_abs,
              "Expected 0. Forecast ties by construction; historical gap = reported cash vs derived.")

    # 3. Revenue reconciliation (constructed vs reported)
    has_segments = spec.disclosure.has_segments
    if has_segments:
        cs.label(row, 1, "Revenue: reported (per filing)")
        rep_row = row
        for p in periods:
            dp = db.get("revenue", p, _CONSOL)
            if dp and dp.value is not None:
                cs.put(rep_row, col_of[p], dp.value, DataType.REPORTED,
                       comment=fw.source_comment(dp))
        row += 1

        def rev_diff(p):
            built = registry.ref("Revenue Build", "revenue", p)
            rep_addr = f"{get_column_letter(col_of[p])}{rep_row}"
            dp = db.get("revenue", p, _CONSOL)
            if built and dp and dp.value is not None:
                return f"={built}-{rep_addr}"
            return None
        add_check("Revenue: segment sum - reported", rev_diff,
                  lambda p: f"MAX({bs_abs},{bs_rel}*ABS({registry.ref('Revenue Build','revenue',p)}))"
                  if registry.ref('Revenue Build', 'revenue', p) else None,
                  "Segment build should reconcile to reported consolidated revenue.")

    # 4. Debt reconciliation
    if "Debt" in spec.schedules:
        def debt_diff(p):
            closing = registry.ref("Debt", "closing_debt", p)
            std = registry.ref("Balance Sheet", "short_term_debt", p)
            ltd = registry.ref("Balance Sheet", "long_term_debt", p)
            terms = [t for t in (std, ltd) if t]
            if closing and terms:
                return f"={closing}-({'+'.join(terms)})"
            return None
        add_check("Debt schedule closing vs Balance Sheet debt", debt_diff, lambda p: bs_abs,
                  "Expected 0.")

    # Overall banner: count of ERROR statuses across the grid
    cs.label(row, 1, "Overall (count of ERROR statuses)", bold=True)
    if status_cells:
        expr = "+".join(f'({c}="ERROR")' for c in status_cells)
        ws.cell(row=row, column=2, value="=SUMPRODUCT(--(" + expr + "))")
        ws.cell(row=row, column=2).font = styles.base_font(bold=True)

    # Conditional formatting: colour PASS/WARNING/ERROR status text cells across the grid.
    last_col = get_column_letter(1 + len(periods))
    grid = f"{get_column_letter(2)}{hdr+1}:{last_col}{row}"
    anchor = f"B{hdr+1}"
    for token, key in [("PASS", "PASS"), ("WARNING", "WARNING"), ("ERROR", "ERROR")]:
        ws.conditional_formatting.add(grid, FormulaRule(
            formula=[f'{anchor}="{token}"'], fill=styles.check_fill(key)))
