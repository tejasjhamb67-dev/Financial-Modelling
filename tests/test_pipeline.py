"""End-to-end and unit tests for the Financial Model Construction Agent."""

from __future__ import annotations

import os

import pytest

from modules.disclosure_mapping.mapper import map_disclosures
from modules.ingestion.loader import load_package
from modules.model_spec import build_spec, order_historical_periods
from modules.normalization.normalizer import normalize
from modules.normalization import restatements as restate
from modules.pipeline import run
from modules.schemas import DataType, SourceDatabase, ValidationStatus
from modules.validation.gates import gate1_data_validation, gate2_model_spec

HERE = os.path.dirname(os.path.abspath(__file__))
PKG = os.path.join(HERE, "fixtures", "synthetic_company.yaml")


@pytest.fixture(scope="module")
def db():
    return normalize(load_package(PKG))


def test_ingestion_loads_datapoints(db):
    assert db.company == "Nova Specialty Chemicals Ltd"
    assert db.value("revenue", "FY26") == 7300
    assert len(db.documents) == 2


def test_missing_data_not_fabricated(db):
    dp = db.get("utilisation_pct", "FY22")
    assert dp is not None
    assert dp.value is None                 # never invented
    assert dp.status == ValidationStatus.NOT_FOUND


def test_restatement_recorded(db):
    restate.apply(db)
    assert any(r.metric == "revenue" and r.period == "FY24" for r in db.restatements)


def test_chronology_oldest_to_newest():
    assert order_historical_periods(["FY26", "FY22", "FY24"]) == ["FY22", "FY24", "FY26"]


def test_spec_three_forecast_years(db):
    spec = build_spec(db, map_disclosures(db))
    assert len(spec.forecast_periods) == 3
    assert spec.forecast_periods == ["FY27E", "FY28E", "FY29E"]
    # forecast follows historical
    assert spec.historical_periods[-1] == "FY26"
    # balance-sheet totals present (recursive resolvability)
    assert "total_assets" in spec.balance_sheet_lines
    assert "total_equity_and_liabilities" in spec.balance_sheet_lines


def test_gates_pass(db):
    g1 = gate1_data_validation(db)
    assert g1.passed
    spec = build_spec(db, map_disclosures(db))
    g2 = gate2_model_spec(spec)
    assert g2.passed


def test_disclosure_mapping(db):
    dm = map_disclosures(db)
    assert dm.has_segments
    assert dm.has_working_capital
    assert dm.has_debt
    assert dm.has_capex_or_dna


def test_full_pipeline_produces_workbook(tmp_path):
    result = run(package=PKG, output_dir=str(tmp_path))
    assert result.success
    assert os.path.exists(result.workbook_path)
    assert os.path.exists(result.build_report_path)
    assert os.path.exists(result.diligence_path)
    assert result.qc_result.passed


def test_reported_financials_sheet(tmp_path):
    from openpyxl import load_workbook
    result = run(package=PKG, output_dir=str(tmp_path))
    wb = load_workbook(result.workbook_path)
    assert "Reported Financials" in wb.sheetnames
    # must sit right before Assumptions (verbatim raw layer)
    assert wb.sheetnames.index("Reported Financials") < wb.sheetnames.index("Assumptions")


def test_source_db_roundtrip(tmp_path, db):
    p = str(tmp_path / "reg.json")
    db.save(p)
    db2 = SourceDatabase.load(p)
    assert db2.value("revenue", "FY26") == 7300
    assert db2.company == db.company


# ---------------------------------------------------------------------------
# Numeric integration test - evaluate the real workbook formulas.
# ---------------------------------------------------------------------------

def _evaluate(path):
    formulas = pytest.importorskip("formulas")
    import warnings
    warnings.filterwarnings("ignore")
    sol = formulas.ExcelModel().loads(path).finish().calculate()

    def val(sheet, cell):
        for k, v in sol.items():
            if f"]{sheet.upper()}'!{cell}" in k.upper():
                try:
                    return v.value[0, 0]
                except Exception:
                    return v.value
        return None
    return val


def test_model_reconciles_numerically(tmp_path):
    result = run(package=PKG, output_dir=str(tmp_path))
    val = _evaluate(result.workbook_path)
    cols = "BCDEFGHI"  # 5 historical + 3 forecast

    # Balance-sheet identity and cash-flow tie hold for every period (Checks sheet
    # row 5 = BS difference, row 9 = CF closing-cash difference).
    for c in cols:
        bs = val("Checks", f"{c}5")
        assert bs is not None and abs(float(bs)) < 1.0, f"BS check {c}={bs}"
        cf = val("Checks", f"{c}9")
        assert cf is not None and abs(float(cf)) < 1.0, f"CF check {c}={cf}"

    # Income statement constructs PAT from reported leaves
    assert abs(float(val("Income Statement", "F16")) - 575.0) < 1.0

    # DCF produces a positive implied value per share
    vps = val("Valuation", "B25")
    assert vps is not None and float(vps) > 0
